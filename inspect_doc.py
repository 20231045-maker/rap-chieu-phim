import docx

doc = docx.Document("BTL_DACN_CNPM.docx")
print("Total Paragraphs:", len(doc.paragraphs))
print("Total Tables:", len(doc.tables))

headings = []
for i, p in enumerate(doc.paragraphs):
    if p.text.startswith("CHƯƠNG") or p.text in ["LỜI CẢM ƠN", "LỜI MỞ ĐẦU", "MỤC LỤC", "DANH MỤC CÁC TỪ VIẾT TẮT", "DANH MỤC HÌNH VẼ", "DANH MỤC BẢNG BIỂU", "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "TÀI LIỆU THAM KHẢO", "PHỤ LỤC"]:
        headings.append((i, p.text))

print("Headings breakdown:")
for idx, (p_idx, text) in enumerate(headings):
    next_idx = headings[idx+1][0] if idx+1 < len(headings) else len(doc.paragraphs)
    print(f"  {text}: paragraphs {p_idx} to {next_idx} (count: {next_idx - p_idx})")
