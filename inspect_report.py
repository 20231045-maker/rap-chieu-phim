import os
import docx

fpath = 'BTL_DACN_CNPM.docx'
sz = os.path.getsize(fpath) / (1024 * 1024)
doc = docx.Document(fpath)

print(f"File: {fpath}")
print(f"Size: {sz:.2f} MB")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")

total_cells = sum(len(t.rows)*len(t.columns) for t in doc.tables)
print(f"Total Table Cells: {total_cells}")

num_imgs = 0
for r in doc.part.rels.values():
    if 'image' in r.target_ref:
        num_imgs += 1
print(f"Total Embedded Images: {num_imgs}")

text_words = sum(len(p.text.split()) for p in doc.paragraphs)
table_words = 0
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            table_words += len(cell.text.split())

print(f"Paragraph Words: {text_words}")
print(f"Table Words: {table_words}")
print(f"Total Words: {text_words + table_words}")

page_breaks = 0
for p in doc.paragraphs:
    if 'w:br w:type="page"' in p._p.xml or 'w:type="page"' in p._p.xml:
        page_breaks += 1
print(f"Explicit Page Breaks: {page_breaks}")

# Also check headings
headings = [p.text for p in doc.paragraphs if p.runs and p.runs[0].bold and len(p.text) < 120 and (p.text.startswith('CHƯƠNG') or any(p.text.startswith(x) for x in ['1.', '2.', '3.', '4.', '5.', 'MỤC LỤC', 'LỜI', 'DANH MỤC', 'KẾT LUẬN', 'TÀI LIỆU', 'PHỤ LỤC']))]
print(f"Headings detected: {len(headings)}")
for h in headings[:25]:
    print("  *", h)
