# ==============================================================================
# SCRIPT XÂY DỰNG BÁO CÁO TOÀN DIỆN BTL_DACN_CNPM.DOCX
# Học phần: ĐỒ ÁN CHUYÊN NGÀNH CÔNG NGHỆ PHẦN MỀM
# Đề tài: HỆ THỐNG ĐẶT VÉ VÀ QUẢN LÝ SUẤT CHIẾU RẠP CHIẾU PHIM (CINESTAR CINEMA)
# Nhóm 10: Nguyễn Mạnh Quyết (20231053) & Hà Thái Bảo (20231045)
# ==============================================================================
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Import các phân hệ chương
from report_sections.cover_and_intro import add_cover_and_intro
from report_sections.chapter1 import add_chapter_1
from report_sections.chapter2 import add_chapter_2
from report_sections.chapter3 import add_chapter_3
from report_sections.chapter4 import add_chapter_4
from report_sections.chapter5 import add_chapter_5
from report_sections.conclusion_and_refs import add_conclusion_and_refs

def build_complete_document():
    print("=================================================================")
    print("🚀 BẮT ĐẦU XÂY DỰNG BÁO CÁO BTL_DACN_CNPM.DOCX (FULL 5 CHƯƠNG)")
    print("=================================================================")

    doc = docx.Document()

    # 1. THIẾT LẬP ĐỊNH DẠNG TRANG CHUẨN THEO CÁCH TRÌNH BÀY .DOCX
    # Lề: Trái 3.0 cm, Phải 2.0 cm, Trên 2.5 cm, Dưới 2.0 cm
    # Khổ giấy A4: 21.0 cm x 29.7 cm
    section = doc.sections[0]
    section.page_width = Inches(8.27)    # 21.0 cm
    section.page_height = Inches(11.69)  # 29.7 cm
    section.top_margin = Inches(0.984)   # 2.5 cm
    section.bottom_margin = Inches(0.787)# 2.0 cm
    section.left_margin = Inches(1.181)  # 3.0 cm
    section.right_margin = Inches(0.787) # 2.0 cm

    # Thiết lập Style Normal mặc định
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(2.0)
    style_normal.paragraph_format.space_before = Pt(0)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Helper: Đánh số trang ở Footer (căn giữa)
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run()
    footer_run.font.name = 'Times New Roman'
    footer_run.font.size = Pt(11)
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    footer_run._r.append(fldChar1)
    footer_run._r.append(instrText)
    footer_run._r.append(fldChar2)
    footer_run._r.append(fldChar3)

    # ---------------------------------------------------------
    # CÁC HÀM TIỆN ÍCH ĐỊNH DẠNG
    # ---------------------------------------------------------
    def add_p(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=2.0, font_size=13):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.italic = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_bullet(text, bold_prefix="", level=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        
        bullet_run = p.add_run("• ")
        bullet_run.bold = True
        bullet_run.font.name = 'Times New Roman'
        bullet_run.font.size = Pt(13)
        
        if bold_prefix:
            prefix_run = p.add_run(bold_prefix + " ")
            prefix_run.bold = True
            prefix_run.font.name = 'Times New Roman'
            prefix_run.font.size = Pt(13)
            
        content_run = p.add_run(text)
        content_run.font.name = 'Times New Roman'
        content_run.font.size = Pt(13)
        return p

    def add_figure(img_rel_path, caption, source="Nhóm 10 - Tác giả tự xây dựng", width_inches=4.2):
        img_path = os.path.join(os.getcwd(), img_rel_path)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(1)
            run = p_img.add_run()
            run.add_picture(img_path, width=Inches(width_inches))

            # Tên hình: phía DƯỚI hình, căn lề giữa, 13pt In Đậm
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(1)
            p_cap.paragraph_format.space_after = Pt(1)
            p_cap.paragraph_format.keep_with_next = True
            run_cap = p_cap.add_run(caption)
            run_cap.bold = True
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(13)

            # Nguồn trích dẫn: phía Dưới, căn lề Phải, In nghiêng
            p_src = doc.add_paragraph()
            p_src.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_src.paragraph_format.space_before = Pt(0)
            p_src.paragraph_format.space_after = Pt(2)
            run_src = p_src.add_run(f"(Nguồn: {source})")
            run_src.italic = True
            run_src.font.name = 'Times New Roman'
            run_src.font.size = Pt(12)
        else:
            print("⚠️ CẢNH BÁO: Không tìm thấy tệp ảnh:", img_path)

    def add_table_custom(title, headers, rows, col_widths=None, align_cols=None, source="Nhóm 10 - Tác giả tổng hợp"):
        # Tên bảng: phía TRÊN bảng, In Đậm
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(5)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        run_cap = p_cap.add_run(title)
        run_cap.bold = True
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(13)

        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Định dạng Header
        hdr_cells = table.rows[0].cells
        for idx, h_text in enumerate(headers):
            hdr_cells[idx].text = h_text
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)
            shading = parse_xml(r'<w:shd {} w:fill="E8E8E8"/>'.format(nsdecls('w')))
            hdr_cells[idx]._tc.get_or_add_tcPr().append(shading)

        # Định dạng Rows
        for r_idx, row_data in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = str(cell_value)
                p = row_cells[c_idx].paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                
                if align_cols and c_idx < len(align_cols):
                    p.alignment = align_cols[c_idx]
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)

        # Borders & Widths
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))
            for c_idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if col_widths and c_idx < len(col_widths):
                    cell.width = Inches(col_widths[c_idx])
                tcPr = cell._tc.get_or_add_tcPr()
                borders = parse_xml(r'''
                    <w:tcBorders {} >
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
                        <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
                        <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
                    </w:tcBorders>
                '''.format(nsdecls('w')))
                tcPr.append(borders)

        # Trích nguồn dưới bảng
        p_src = doc.add_paragraph()
        p_src.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_src.paragraph_format.space_before = Pt(2)
        p_src.paragraph_format.space_after = Pt(4)
        run_src = p_src.add_run(f"(Nguồn: {source})")
        run_src.italic = True
        run_src.font.name = 'Times New Roman'
        run_src.font.size = Pt(12)

    def add_code_snippet(code_text, lang="SQL / JavaScript"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.right_indent = Inches(0.2)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        shd = parse_xml(r'<w:shd {} w:fill="F4F4F4"/>'.format(nsdecls('w')))
        p._p.get_or_add_pPr().append(shd)

    # 2. XÂY DỰNG TUẦN TỰ TỪNG PHẦN THEO KẾ HOẠCH
    print("1. Đang xây dựng Trang bìa, Lời mở đầu, Mục lục, Danh mục...")
    add_cover_and_intro(doc, add_p, add_heading_1, add_heading_2, add_heading_3, add_bullet, add_figure, add_table_custom)

    print("2. Đang xây dựng CHƯƠNG I: Tổng quan về đề tài...")
    add_chapter_1(doc, add_p, add_heading_1, add_heading_2, add_heading_3, add_bullet, add_figure, add_table_custom)

    print("3. Đang xây dựng CHƯƠNG II: Tổng hợp, phân tích và đánh giá thực tế...")
    add_chapter_2(doc, add_p, add_heading_1, add_heading_2, add_heading_3, add_bullet, add_figure, add_table_custom)

    print("4. Đang xây dựng CHƯƠNG III: Cơ sở lý thuyết...")
    add_chapter_3(doc, add_p, add_heading_1, add_heading_2, add_heading_3, add_bullet, add_figure, add_table_custom, add_code_snippet)

    print("5. Đang xây dựng CHƯƠNG IV: Phân tích và thiết kế hệ thống...")
    add_chapter_4(doc, add_p, add_heading_1, add_heading_2, add_heading_3, add_bullet, add_figure, add_table_custom, add_code_snippet)

    print("6. Đang xây dựng CHƯƠNG V: Cài đặt và triển khai...")
    add_chapter_5(doc, add_p, add_heading_1, add_heading_2, add_heading_3, add_bullet, add_figure, add_table_custom, add_code_snippet)

    print("7. Đang xây dựng Kết luận, Tài liệu tham khảo, Phụ lục...")
    add_conclusion_and_refs(doc, add_p, add_heading_1, add_heading_2, add_heading_3, add_bullet, add_figure, add_table_custom, add_code_snippet)

    output_path = os.path.join(os.getcwd(), "BTL_DACN_CNPM.docx")
    doc.save(output_path)
    print("=================================================================")
    print(f"🎉 BÁO CÁO ĐÃ ĐƯỢC XUẤT THÀNH CÔNG TẠI: {output_path}")
    print("=================================================================")

if __name__ == "__main__":
    build_complete_document()
